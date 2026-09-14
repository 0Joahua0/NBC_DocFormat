"""Line-spacing mode regression tests."""

from copy import deepcopy

from docx import Document
from docx.enum.text import WD_LINE_SPACING

from NBC_DocFormat import DEFAULT_CUSTOM_SETTINGS
from scripts.formatter import PRESETS, format_document


def _make_source(path):
    document = Document()
    document.add_paragraph("关于行距设置的通知")
    document.add_paragraph("正文段落用于验证行距。")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格内容"
    document.save(path)


def test_every_custom_format_element_declares_a_line_spacing_mode():
    for key in (
        "title", "recipient", "heading1", "heading2", "heading3", "heading4",
        "body", "signature", "date", "attachment", "closing", "table",
    ):
        assert DEFAULT_CUSTOM_SETTINGS[key]["line_spacing_type"] in {"exact", "multiple"}


def test_body_supports_multiple_line_spacing(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    _make_source(source)

    settings = deepcopy(PRESETS["official"])
    settings["page_number"] = False
    settings["body"].update({"line_spacing_type": "multiple", "line_spacing": 1.23})

    format_document(str(source), str(output), preset_name="custom", custom_settings=settings)

    paragraph = next(p for p in Document(output).paragraphs if "正文段落" in p.text)
    assert paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.MULTIPLE
    assert abs(float(paragraph.paragraph_format.line_spacing) - 1.23) < 0.02


def test_body_keeps_exact_point_line_spacing(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    _make_source(source)

    settings = deepcopy(PRESETS["official"])
    settings["page_number"] = False
    settings["body"].update({"line_spacing_type": "exact", "line_spacing": 31.5})

    format_document(str(source), str(output), preset_name="custom", custom_settings=settings)

    paragraph = next(p for p in Document(output).paragraphs if "正文段落" in p.text)
    assert paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
    assert abs(paragraph.paragraph_format.line_spacing.pt - 31.5) < 0.02


def test_table_explicit_spacing_overrides_single_spacing_default(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    _make_source(source)

    settings = deepcopy(PRESETS["official"])
    settings["page_number"] = False
    settings.setdefault("table", {}).update({
        "line_spacing_type": "multiple",
        "line_spacing": 1.23,
    })

    format_document(str(source), str(output), preset_name="custom", custom_settings=settings)

    paragraph = Document(output).tables[0].cell(0, 0).paragraphs[0]
    assert paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.MULTIPLE
    assert abs(float(paragraph.paragraph_format.line_spacing) - 1.23) < 0.02
