"""Exercise real mapped Tk windows, including the Linux/X11 grab lifecycle."""

import tkinter as tk
import json

import pytest
from docx import Document

import NBC_DocFormat as app


@pytest.fixture
def root(tmp_path, monkeypatch):
    monkeypatch.setattr(app, 'CONFIG_FILE', tmp_path / 'settings.json')
    try:
        window = tk.Tk(className='NBC_DocFormat')
    except tk.TclError as error:
        pytest.skip(f'Tk display is unavailable: {error}')
    errors = []
    window.report_callback_exception = lambda *args: errors.append(args)
    window.update()
    yield window
    window.destroy()
    assert not errors


@pytest.mark.parametrize('kind', ['paste', 'custom', *app.BUILTIN_PRESET_IDS])
def test_dialog_has_widgets_and_is_visible_before_grabbing(root, monkeypatch, kind):
    original_grab = tk.Toplevel.grab_set
    grabbed = []

    def require_visible(dialog):
        assert dialog.winfo_viewable(), 'X11 cannot grab an unmapped dialog'
        assert dialog.winfo_children(), 'Do not grab before building the UI'
        grabbed.append(dialog)
        original_grab(dialog)

    monkeypatch.setattr(tk.Toplevel, 'grab_set', require_visible)
    if kind == 'paste':
        dialog = app.PasteTextDialog(root)
    else:
        dialog = app.CustomSettingsDialog(root, preset_id=None if kind == 'custom' else kind)
    root.update()
    assert grabbed == [dialog]
    assert root.grab_current() == dialog
    dialog.destroy()
    root.update()
    assert root.grab_current() is None


def test_rejected_grab_does_not_leave_a_blank_dialog(root, monkeypatch):
    def reject_grab(dialog):
        raise tk.TclError('grab failed: window not viewable')

    monkeypatch.setattr(tk.Toplevel, 'grab_set', reject_grab)
    dialog = app.PasteTextDialog(root)
    root.update()
    assert dialog.text_widget.winfo_viewable()
    assert dialog.title_var.get() == '新建文档'
    dialog._on_close()
    assert root.grab_current() is None


def test_close_before_visibility_does_not_block_or_grab_later(root):
    root.withdraw()
    dialog = app.PasteTextDialog(root)
    dialog._on_close()
    root.deiconify()
    root.update()
    assert root.grab_current() is None


def test_clipboard_text_can_be_generated_and_formatted(root, tmp_path):
    output = tmp_path / '测试通知.docx'
    received = []

    def generate(title, text, output_path, is_markdown):
        received.append((title, text, output_path, is_markdown))
        source = tmp_path / 'input.docx'
        app._create_docx_from_text(title, text, source)
        app.format_document(str(source), output_path, preset_name='nbc')

    dialog = app.PasteTextDialog(root, on_generate=generate)
    root.update()
    dialog.title_var.set('测试通知')
    dialog.save_path_var.set(str(tmp_path))
    root.clipboard_clear()
    root.clipboard_append('这是粘贴的正文。\n\n这是第二段正文。')
    dialog.text_widget.event_generate('<<Paste>>')
    root.update()
    dialog._generate()
    assert received == [('测试通知', '这是粘贴的正文。\n\n这是第二段正文。', str(output), False)]
    assert [p.text for p in Document(output).paragraphs if p.text.strip()] == [
        '测试通知', '这是粘贴的正文。', '这是第二段正文。',
    ]


def test_application_icon_is_loaded_and_retained(root):
    app._set_window_icon(root)
    assert root._nbc_icon.width() > 0
    assert root._nbc_icon.height() > 0
    assert root.winfo_class() == 'Nbc_docformat'


def test_main_window_and_its_dialogs_render(root):
    app.DocFormatApp(root)
    root.update()
    for dialog_class in (app.PasteTextDialog, app.CustomSettingsDialog):
        dialog = app._open_dialog(root, dialog_class)
        root.update()
        assert dialog is not None
        assert dialog.winfo_viewable()
        assert dialog.winfo_children()
        dialog.destroy()
        root.update()


def test_docx_generation_preserves_supplementary_unicode(tmp_path):
    output = tmp_path / 'unicode.docx'
    title = '𠮷野的通知 📋'
    body = '文档正文保留 emoji 🧾 和扩展汉字𠮷。'
    app._create_docx_from_text(title, body, output)
    assert [p.text for p in Document(output).paragraphs if p.text] == [title, body]


@pytest.mark.parametrize('dialog_class, builder', [
    (app.PasteTextDialog, '_build_ui'),
    (app.CustomSettingsDialog, '_create_widgets'),
])
def test_failed_dialog_build_reports_stage_and_removes_only_failed_window(
        root, monkeypatch, dialog_class, builder):
    other = tk.Toplevel(root)
    messages = []
    monkeypatch.setattr(app.messagebox, 'showerror', lambda *args, **kwargs: messages.append(args))

    def broken_builder(dialog):
        tk.Frame(dialog).pack()
        raise tk.TclError('simulated Linux widget creation failure')

    monkeypatch.setattr(dialog_class, builder, broken_builder)
    assert app._open_dialog(root, dialog_class) is None
    root.update()
    assert root.winfo_children() == [other]
    assert root.grab_current() is None
    records = [json.loads(line) for line in
               (app.CONFIG_FILE.parent / 'ui_diagnostics.log').read_text(encoding='utf-8').splitlines()]
    assert any(record['event'] == 'ui_environment' and record['tk'] for record in records)
    error = next(record for record in records if record['event'] == 'ui_error')
    assert 'broken_builder' in error['traceback']
    assert 'simulated Linux widget creation failure' in error['traceback']
    assert dialog_class.__name__ in error['context']
    assert messages and 'ui_diagnostics.log' in messages[0][1]


def test_visible_dialog_logs_widget_state_without_input_text(root):
    dialog = app._open_dialog(root, app.PasteTextDialog)
    dialog.text_widget.insert('end', 'PRIVATE_DOCUMENT_CONTENT_12345')
    root.after(650, root.quit)
    root.mainloop()
    log = (app.CONFIG_FILE.parent / 'ui_diagnostics.log').read_text(encoding='utf-8')
    assert 'PRIVATE_DOCUMENT_CONTENT_12345' not in log
    records = [json.loads(line) for line in log.splitlines()]
    snapshots = [record for record in records if record['event'] == 'dialog_state']
    assert len(snapshots) >= 2
    assert snapshots[-1]['viewable']
    assert snapshots[-1]['widget_counts']['Text'] == 1
    assert snapshots[-1]['content']['text_widget']['viewable']
    assert snapshots[-1]['viewable_widgets'] > 0


def test_close_cancels_pending_diagnostic_snapshot(root):
    before = set(root.tk.call('after', 'info'))
    dialog = app._open_dialog(root, app.PasteTextDialog)
    assert set(root.tk.call('after', 'info')) - before
    dialog.destroy()
    assert not (set(root.tk.call('after', 'info')) - before)
    root.update()


def test_unwritable_diagnostic_log_does_not_prevent_dialog_opening(root, tmp_path, monkeypatch):
    blocked = tmp_path / 'not-a-directory'
    blocked.write_text('occupied', encoding='utf-8')
    monkeypatch.setattr(app, 'CONFIG_FILE', blocked / 'settings.json')
    dialog = app._open_dialog(root, app.PasteTextDialog)
    root.update()
    assert dialog.text_widget.winfo_viewable()


def test_diagnostic_log_rotates_at_size_limit(tmp_path, monkeypatch):
    monkeypatch.setattr(app, 'CONFIG_FILE', tmp_path / 'settings.json')
    log = tmp_path / 'ui_diagnostics.log'
    log.write_bytes(b'x' * (1024 * 1024))
    assert app._write_ui_diagnostic('rotation_check') == log
    assert log.with_suffix('.log.1').stat().st_size == 1024 * 1024
    assert json.loads(log.read_text(encoding='utf-8'))['event'] == 'rotation_check'
