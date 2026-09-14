"""Middle-dot punctuation font normalization."""

from docx import Document
from docx.oxml.ns import qn

from scripts.analyzer import analyze_punctuation
from scripts.formatter import PRESETS, format_document
from scripts.punctuation import MIDDLE_DOT_FONT_CN, fix_text, process_document


def _run_font_values(run):
    rfonts = run._r.rPr.find(qn("w:rFonts"))
    assert rfonts is not None
    return {attr: rfonts.get(qn("w:" + attr)) for attr in ("eastAsia", "ascii", "hAnsi", "cs")}


def test_middle_dot_is_split_and_set_to_required_font(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"

    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("亲清直通车·企需关应")
    run.font.name = "Times New Roman"
    document.save(source)

    assert analyze_punctuation(Document(source))[0]["type"] == "间隔号字体不规范"

    process_document(str(source), str(output))

    fixed = Document(output).paragraphs[0]
    assert fixed.text == "亲清直通车·企需关应"
    middle_dot_run = next(run for run in fixed.runs if run.text == "·")
    assert all(value == MIDDLE_DOT_FONT_CN for value in _run_font_values(middle_dot_run).values())
    assert not analyze_punctuation(Document(output))


def test_ratio_colon_is_not_changed_to_math_ratio_symbol():
    assert fix_text("示例：“4：6”，不得出现“4：6”。") == "示例：“4：6”，不得出现“4：6”。"


def test_formatter_preserves_required_middle_dot_font(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"

    document = Document()
    document.add_paragraph("关于间隔号字体的通知")
    document.add_paragraph("请使用亲清直通车·企需关应。")
    document.save(source)

    settings = PRESETS["official"].copy()
    settings["page_number"] = False
    format_document(str(source), str(output), preset_name="custom", custom_settings=settings)

    paragraph = next(p for p in Document(output).paragraphs if "亲清直通车" in p.text)
    middle_dot_run = next(run for run in paragraph.runs if run.text == "·")
    assert all(value == MIDDLE_DOT_FONT_CN for value in _run_font_values(middle_dot_run).values())
