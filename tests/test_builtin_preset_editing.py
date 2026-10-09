"""Built-in preset editing, persistence and document output."""

from copy import deepcopy
from types import SimpleNamespace
import tkinter as tk

import pytest
from docx import Document

import NBC_DocFormat as app


@pytest.fixture(scope='module')
def tk_root():
    try:
        root = tk.Tk()
    except tk.TclError as error:
        pytest.skip(f'Tk display is unavailable: {error}')
    root.withdraw()
    yield root
    root.destroy()


@pytest.fixture
def editor_root(tk_root, tmp_path, monkeypatch):
    monkeypatch.setattr(app, 'CONFIG_FILE', tmp_path / 'settings.json')
    monkeypatch.setattr(app, '_fit_dialog_to_screen', lambda dialog, *a, **kw: dialog.withdraw())
    yield tk_root
    for child in list(tk_root.winfo_children()):
        child.destroy()


def open_editor(root, preset_id):
    dialog = app.CustomSettingsDialog(root, preset_id=preset_id)
    root.update_idletasks()
    return dialog


@pytest.mark.parametrize('preset_id', app.BUILTIN_PRESET_IDS)
def test_save_without_edits_preserves_exact_builtin_settings(editor_root, preset_id):
    original = deepcopy(app.PRESETS)
    dialog = open_editor(editor_root, preset_id)
    dialog._save()
    config = app.load_custom_settings()
    assert not config.get('builtin_overrides')
    assert app.get_format_settings(preset_id) == original[preset_id]
    assert app.PRESETS == original


@pytest.mark.parametrize('preset_id', app.BUILTIN_PRESET_IDS)
def test_edit_persists_independently_and_reset_restores_defaults(editor_root, preset_id):
    config = app.load_custom_settings()
    app.save_custom_settings(config)
    user_preset = deepcopy(app.get_active_user_preset(config))
    dialog = open_editor(editor_root, preset_id)
    dialog.margin_vars['left'].set('1.7')
    dialog.h2_font_var.set('测试字体')
    dialog._save()

    reopened = open_editor(editor_root, preset_id)
    assert reopened.margin_vars['left'].get() == '1.7'
    assert reopened.h2_font_var.get() == '测试字体'
    expected = deepcopy(app.PRESETS[preset_id])
    expected['page']['left'] = 1.7
    expected['heading2']['font_cn'] = '测试字体'
    assert app.get_format_settings(preset_id) == expected
    assert app.get_active_user_preset(app.load_custom_settings()) == user_preset
    for other_id in app.BUILTIN_PRESET_IDS:
        if other_id != preset_id:
            assert app.get_format_settings(other_id) == app.PRESETS[other_id]

    reopened._reset_defaults()
    reopened._save()
    assert app.get_format_settings(preset_id) == app.PRESETS[preset_id]
    assert preset_id not in app.load_custom_settings()['builtin_overrides']


def test_advanced_edit_can_match_body_font_without_losing_other_nbc_fields(editor_root):
    dialog = open_editor(editor_root, 'nbc')
    dialog._adv_vars['date']['font'].set(app.PRESETS['nbc']['body']['font_cn'])
    dialog._adv_vars['date']['bold'].set(False)
    dialog._save()
    expected = deepcopy(app.PRESETS['nbc'])
    expected['date']['font_cn'] = expected['body']['font_cn']
    expected['date']['bold'] = False
    assert app.get_format_settings('nbc') == expected


def test_cancel_discards_edits(editor_root, monkeypatch):
    dialog = open_editor(editor_root, 'nbc')
    dialog.margin_vars['left'].set('1.2')
    monkeypatch.setattr(app.messagebox, 'askyesnocancel', lambda *a, **kw: False)
    dialog._on_close()
    assert not app.CONFIG_FILE.exists()
    assert app.get_format_settings('nbc') == app.PRESETS['nbc']


def test_failed_save_keeps_editor_open_and_does_not_call_success(editor_root, monkeypatch):
    dialog = open_editor(editor_root, 'nbc')
    saved, errors = [], []
    dialog.on_save = saved.append
    monkeypatch.setattr(app, 'save_custom_settings', lambda config: None)
    monkeypatch.setattr(app.messagebox, 'showerror', lambda *a, **kw: errors.append(a))
    dialog._save()
    assert dialog.winfo_exists()
    assert errors and not saved


def test_gui_format_uses_saved_builtin_settings(editor_root, tmp_path):
    dialog = open_editor(editor_root, 'nbc')
    dialog.margin_vars['left'].set('1.7')
    dialog.body_size_var.set('小三(15pt)')
    dialog._save()
    source, output = tmp_path / 'source.docx', tmp_path / 'output.docx'
    document = Document()
    document.add_paragraph('关于测试的通知')
    document.add_paragraph('这是用于验证预设修改生效的正文。')
    document.save(source)
    controller = SimpleNamespace(
        preset=tk.StringVar(editor_root, value='nbc'),
        log_panel=SimpleNamespace(log=lambda *a: None),
    )
    app.DocFormatApp._run_format(controller, str(source), str(output), progress_callback=lambda *a: None)
    formatted = Document(output)
    assert formatted.sections[0].left_margin.cm == pytest.approx(1.7, abs=0.01)
    assert all(run.font.size.pt == 15 for run in formatted.paragraphs[1].runs if run.text.strip())


def test_card_edit_button_opens_its_preset_and_respects_disabled_state(editor_root):
    selected = tk.StringVar(editor_root, value='official')
    opened = []
    cards = [app.PresetCard(editor_root, preset_id, preset_id, selected,
                           edit_command=lambda key=preset_id: opened.append(key))
             for preset_id in app.BUILTIN_PRESET_IDS]
    for card in cards:
        card.edit_button.invoke()
        card.set_enabled(False)
        card.edit_button.invoke()
    assert opened == list(app.BUILTIN_PRESET_IDS)
    assert selected.get() == 'official'
