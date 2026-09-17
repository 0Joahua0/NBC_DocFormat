"""NBC preset values and output formatting."""

import re
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from scripts.formatter import PRESETS, format_document
from scripts.punctuation import MIDDLE_DOT_FONT_CN, iter_paragraph_runs


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
FOOTNOTES_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
)


def _east_asian_font(paragraph):
    rfonts = paragraph.runs[0]._r.rPr.find(qn("w:rFonts"))
    assert rfonts is not None
    return rfonts.get(qn("w:eastAsia"))


def _run_font_values(run):
    rfonts = run._r.rPr.find(qn("w:rFonts"))
    assert rfonts is not None
    return {
        attribute: rfonts.get(qn(f"w:{attribute}"))
        for attribute in ("eastAsia", "ascii", "hAnsi", "cs")
    }


def _assert_contextual_punctuation_fonts(paragraph, expected_chinese_font):
    quote_chars = set("“”‘’")
    runs = iter_paragraph_runs(paragraph)
    quote_runs = [run for run in runs if any(char in run.text for char in quote_chars)]
    assert quote_runs
    for run in quote_runs:
        assert set(run.text) <= quote_chars
        assert all(value == expected_chinese_font for value in _run_font_values(run).values())
        rfonts = run._r.rPr.find(qn("w:rFonts"))
        assert rfonts.get(qn("w:hint")) == "eastAsia"

    latin_digit_runs = [run for run in runs if re.search(r"[A-Za-z0-9]", run.text)]
    assert latin_digit_runs
    for run in latin_digit_runs:
        assert not any(char in run.text for char in quote_chars)
        fonts = _run_font_values(run)
        assert fonts["ascii"] == "Times New Roman"
        assert fonts["hAnsi"] == "Times New Roman"
        assert fonts["cs"] == "Times New Roman"

    for run in runs:
        if "·" not in run.text:
            continue
        assert run.text == "·"
        assert all(value == MIDDLE_DOT_FONT_CN for value in _run_font_values(run).values())
        rfonts = run._r.rPr.find(qn("w:rFonts"))
        assert rfonts.get(qn("w:hint")) == "eastAsia"


def _append_wrapped_run(paragraph, text):
    hyperlink = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_test_footnote(document):
    footnotes_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="{W_NS}">
  <w:footnote w:id="1">
    <w:p>
      <w:pPr><w:spacing w:line="560" w:lineRule="exact"/></w:pPr>
      <w:r><w:footnoteRef/></w:r>
      <w:r><w:t>脚注“AI2026·测试”</w:t></w:r>
    </w:p>
  </w:footnote>
</w:footnotes>""".encode("utf-8")
    part = Part(
        PackURI("/word/footnotes.xml"),
        FOOTNOTES_CONTENT_TYPE,
        footnotes_xml,
        document.part.package,
    )
    document.part.relate_to(part, RT.FOOTNOTES)

    paragraph = document.add_paragraph("含脚注的正文")
    reference = OxmlElement("w:footnoteReference")
    reference.set(qn("w:id"), "1")
    paragraph.add_run()._r.append(reference)


def test_nbc_replaces_legal_builtin_preset():
    assert "nbc" in PRESETS
    assert "legal" not in PRESETS

    preset = PRESETS["nbc"]
    assert preset["name"] == "NBC格式"
    assert preset["page_number_font"] == "Times New Roman"
    assert preset["page_number_position"] == "center"
    assert preset["page"] == {
        "top": 3.4,
        "bottom": 2.8,
        "left": 2.8,
        "right": 2.8,
        "header_distance": 1.5,
        "footer_distance": 1.75,
    }
    assert preset["body"]["font_cn"] == "方正仿宋_GBK"
    assert preset["body"]["size"] == 18
    assert preset["body"]["line_spacing_type"] == "multiple"
    assert preset["body"]["line_spacing"] == 1.23
    assert preset["heading2"]["font_cn"] == "方正楷体_GBK"
    assert preset["heading2"]["bold"] is True
    assert preset["heading3"]["bold"] is True
    assert preset["date"]["font_cn"] == "方正楷体_GBK"
    assert preset["date"]["bold"] is True
    assert preset["date"]["align"] == "center"
    assert preset["date"]["line_spacing"] == 1.0
    assert preset["footnote"] == {
        "font_cn": "方正仿宋_GBK",
        "font_en": "Times New Roman",
        "size": 12,
        "bold": False,
        "line_spacing_type": "multiple",
        "line_spacing": 1.0,
        "space_before": 0,
        "space_after": 0,
    }


def test_nbc_preset_formats_docx_with_reference_layout(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"

    document = Document()
    document.add_paragraph("关于NBC格式测试的通知")
    document.add_paragraph("（2026年8月13日）")
    document.add_paragraph("一、一级标题")
    document.add_paragraph("（一）二级标题")
    document.add_paragraph("1.三级标题")
    document.add_paragraph("正文段落用于验证NBC格式。")
    document.save(source)

    format_document(
        str(source), str(output),
        preset_name="nbc",
        custom_settings={"page_number": False},
    )

    result = Document(output)
    section = result.sections[0]
    assert abs(section.top_margin.cm - 3.4) < 0.02
    assert abs(section.bottom_margin.cm - 2.8) < 0.02
    assert abs(section.left_margin.cm - 2.8) < 0.02
    assert abs(section.right_margin.cm - 2.8) < 0.02
    assert abs(section.header_distance.cm - 1.5) < 0.02
    assert abs(section.footer_distance.cm - 1.75) < 0.02

    body = next(p for p in result.paragraphs if "正文段落" in p.text)
    assert _east_asian_font(body) == "方正仿宋_GBK"
    assert body.runs[0].font.size.pt == 18
    assert body.paragraph_format.line_spacing_rule == WD_LINE_SPACING.MULTIPLE
    assert abs(float(body.paragraph_format.line_spacing) - 1.23) < 0.02

    heading2 = next(p for p in result.paragraphs if p.text.startswith("（一）"))
    assert _east_asian_font(heading2) == "方正楷体_GBK"
    assert heading2.runs[0].font.bold is True

    date = next(p for p in result.paragraphs if p.text == "（2026年8月13日）")
    assert _east_asian_font(date) == "方正楷体_GBK"
    assert date.runs[0].font.size.pt == 18
    assert date.runs[0].font.bold is True
    assert date.alignment == 1
    assert date.paragraph_format.line_spacing_rule == WD_LINE_SPACING.SINGLE
    assert abs(float(date.paragraph_format.line_spacing) - 1.0) < 0.02


def test_nbc_page_numbers_are_centered_times_new_roman_fields(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"

    document = Document()
    document.add_paragraph("NBC页码格式测试")
    document.save(source)

    format_document(str(source), str(output), preset_name="nbc")

    footer = Document(output).sections[0].footer
    paragraph = footer.paragraphs[0]
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert "— 1 —" in paragraph.text

    for run in paragraph.runs:
        fonts = run._r.rPr.find(qn("w:rFonts"))
        assert fonts is not None
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            assert fonts.get(qn(f"w:{attribute}")) == "Times New Roman"

    field_types = [
        field.get(qn("w:fldCharType"))
        for field in paragraph._p.iter(qn("w:fldChar"))
    ]
    assert field_types == ["begin", "separate", "end"]


def test_nbc_placeholder_date_is_detected_from_reference_template(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"

    document = Document()
    document.add_paragraph("重要文稿材料格式规范")
    document.add_paragraph("（202X年X月X日）")
    document.add_paragraph("材料抬头：")
    document.save(source)

    format_document(
        str(source), str(output),
        preset_name="nbc",
        custom_settings={"page_number": False},
    )

    date = next(p for p in Document(output).paragraphs if p.text == "（202X年X月X日）")
    assert _east_asian_font(date) == "方正楷体_GBK"
    assert date.runs[0].font.bold is True
    assert date.alignment == 1


def test_nbc_formats_real_word_footnotes(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"

    document = Document()
    document.add_paragraph("NBC脚注格式测试")
    _add_test_footnote(document)
    document.save(source)

    format_document(
        str(source), str(output),
        preset_name="nbc",
        custom_settings={"page_number": False},
    )

    with ZipFile(output) as package:
        root = etree.fromstring(package.read("word/footnotes.xml"))

    namespace = {"w": W_NS}
    footnote = root.xpath('./w:footnote[@w:id="1"]', namespaces=namespace)[0]
    spacing = footnote.find(".//w:pPr/w:spacing", namespace)
    assert spacing.get(qn("w:line")) == "240"
    assert spacing.get(qn("w:lineRule")) == "auto"

    text_runs = footnote.xpath('.//w:r[w:t]', namespaces=namespace)
    assert "".join(run.find("w:t", namespace).text for run in text_runs) == "脚注“AI2026·测试”"

    text_run = text_runs[0]
    rfonts = text_run.find("w:rPr/w:rFonts", namespace)
    assert rfonts.get(qn("w:eastAsia")) == "方正仿宋_GBK"
    assert rfonts.get(qn("w:ascii")) == "Times New Roman"
    assert rfonts.get(qn("w:hAnsi")) == "Times New Roman"
    assert text_run.find("w:rPr/w:sz", namespace).get(qn("w:val")) == "24"

    for run in text_runs:
        text = run.find("w:t", namespace).text
        fonts = run.find("w:rPr/w:rFonts", namespace)
        if any(char in text for char in "“”"):
            assert set(text) <= set("“”")
            assert all(
                fonts.get(qn(f"w:{attribute}")) == "方正仿宋_GBK"
                for attribute in ("eastAsia", "ascii", "hAnsi", "cs")
            )
        if re.search(r"[A-Za-z0-9]", text):
            assert fonts.get(qn("w:ascii")) == "Times New Roman"
            assert fonts.get(qn("w:hAnsi")) == "Times New Roman"
        if "·" in text:
            assert text == "·"
            assert all(
                fonts.get(qn(f"w:{attribute}")) == MIDDLE_DOT_FONT_CN
                for attribute in ("eastAsia", "ascii", "hAnsi", "cs")
            )


def test_nbc_quotes_follow_context_while_latin_and_digits_stay_times_new_roman(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"

    document = Document()
    expected_fonts = {
        "关于“聚·AI2026”项目的通知": "方正小标宋_GBK",
        "一、“AI2026”工作": "方正黑体_GBK",
        "（一）‘AI2026’安排": "方正楷体_GBK",
        "正文“AI2026”和‘B2’内容。": "方正仿宋_GBK",
        "一是：“AI2026”取得成效。": "方正仿宋_GBK",
    }
    for text in expected_fonts:
        document.add_paragraph(text)

    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格“AI2026·测试”"
    nested_table = table.cell(0, 0).add_table(rows=1, cols=1)
    nested_table.cell(0, 0).text = "嵌套“AI2026·测试”"

    wrapped_paragraph = document.add_paragraph()
    _append_wrapped_run(wrapped_paragraph, "“AI2026·测试”")
    document.save(source)

    format_document(
        str(source), str(output),
        preset_name="nbc",
        custom_settings={"page_number": False},
    )

    result = Document(output)
    for text, expected_chinese_font in expected_fonts.items():
        paragraph = next(p for p in result.paragraphs if p.text == text)
        _assert_contextual_punctuation_fonts(paragraph, expected_chinese_font)

    table_paragraph = result.tables[0].cell(0, 0).paragraphs[0]
    assert table_paragraph.text == "表格“AI2026·测试”"
    _assert_contextual_punctuation_fonts(table_paragraph, "方正仿宋_GBK")

    nested_paragraph = result.tables[0].cell(0, 0).tables[0].cell(0, 0).paragraphs[0]
    assert nested_paragraph.text == "嵌套“AI2026·测试”"
    _assert_contextual_punctuation_fonts(nested_paragraph, "方正仿宋_GBK")

    wrapped_result = next(
        p for p in result.paragraphs
        if "".join(run.text for run in iter_paragraph_runs(p)) == "“AI2026·测试”"
    )
    _assert_contextual_punctuation_fonts(wrapped_result, "方正仿宋_GBK")
