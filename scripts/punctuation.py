#!/usr/bin/env python3
"""中文标点修复与特殊字符字体归一化。

这个模块同时处理两个层次的问题：

1. **纯文本层**：把英文标点转换为适合中文公文的标点，并对整段引号做成对
   判断。引号必须按整段判断，因为一对引号可能被 Word 拆进不同的 ``run``。
2. **OOXML 层**：只把中文引号和中点拆成独立 ``run``，再设置它们的字体。
   这样同一原始 ``run`` 中的英文和数字仍可继续使用 Times New Roman。

理解本文件前需要知道：Word 的一段文字由若干 ``w:r``（run）组成，一个 run
不仅可能包含 ``w:t`` 文本，还可能包含图片、域、换行等非文本节点。因此这里
不能简单地用 ``paragraph.text = ...`` 重建段落；拆分时必须逐个复制 XML 子节点，
否则会丢失超链接、图片、域代码或修订记录。

对外主要入口：

``process_paragraph``
    修复一个段落的标点、空格和特殊字符字体。
``process_document``
    命令行使用的整篇文档入口。
``normalize_chinese_quote_fonts`` / ``normalize_middle_dot_fonts``
    由格式化器复用的字体归一化函数。
"""

import re
import sys
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.run import Run
try:
    from scripts.east_asian_typography import apply_chinese_line_break_rules
except ModuleNotFoundError:  # Support direct execution from scripts/.
    from east_asian_typography import apply_chinese_line_break_rules

# 中文标点（使用Unicode转义确保正确）
LEFT_DOUBLE_QUOTE = '\u201c'   # " 左双引号
RIGHT_DOUBLE_QUOTE = '\u201d'  # " 右双引号
LEFT_SINGLE_QUOTE = '\u2018'   # ' 左单引号
RIGHT_SINGLE_QUOTE = '\u2019'  # ' 右单引号

# 基本替换映射
REPLACEMENTS = {
    "(": "（",
    ")": "）",
    ":": "：",
    ";": "；",
    "?": "？",
    "!": "！",
}

# 占位符前缀（使用不可能出现在正常文本中的字符序列）
_PLACEHOLDER_PREFIX = "\x02PROT"
MIDDLE_DOT = "·"
MIDDLE_DOT_FONT_CN = "方正仿宋_GBK"
MIDDLE_DOT_ALLOWED_FONTS = {MIDDLE_DOT_FONT_CN, "方正仿宋"}
CHINESE_QUOTE_CHARS = frozenset({
    LEFT_DOUBLE_QUOTE,
    RIGHT_DOUBLE_QUOTE,
    LEFT_SINGLE_QUOTE,
    RIGHT_SINGLE_QUOTE,
})
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def _protect_special_patterns(text):
    """临时隐藏不应中文化的片段，返回 ``(占位文本, 恢复表)``。

    例如时间 ``9:30``、URL 中的冒号和标准号 ``ISO 9001:2015`` 都不能被
    全角化。占位符使用控制字符包围，正常用户文本几乎不可能与其冲突；全部
    标点规则执行完后再由 :func:`_restore_protected` 原样放回。
    """
    protected = []
    counter = [0]

    def _replace_with_placeholder(match):
        placeholder = f"{_PLACEHOLDER_PREFIX}{counter[0]}\x03"
        protected.append((placeholder, match.group()))
        counter[0] += 1
        return placeholder

    result = text

    # 保护 URL（http:// https:// ftp://）
    result = re.sub(r"(?:https?|ftp)://\S+", _replace_with_placeholder, result)

    # 保护邮箱地址
    result = re.sub(r"[\w.+-]+@[\w-]+\.[\w.-]+", _replace_with_placeholder, result)

    # 保护 Windows 文件路径 C:\ D:\
    result = re.sub(r"[A-Za-z]:\\", _replace_with_placeholder, result)

    # 保护标准编号 ISO 9001:2015 等（字母+数字+冒号+数字）
    result = re.sub(r"[A-Za-z]+[\s-]?\d+:\d{2,}", _replace_with_placeholder, result)

    # 保护时间格式 HH:MM 或 HH:MM:SS
    # 注意：不能用 \b，因为 Python3 中中文字符也属于 \w，
    # 导致"上午9:30"中 午 和 9 之间没有 \b 边界
    result = re.sub(r"(?<!\d)(\d{1,2}:\d{2}(?::\d{2})?)(?!\d)", _replace_with_placeholder, result)

    return result, protected


def _restore_protected(text, protected):
    """恢复被保护的内容"""
    result = text
    for placeholder, original in protected:
        result = result.replace(placeholder, original)
    return result


def has_chinese(text):
    """检查是否包含中文"""
    return bool(re.search(r"[\u4e00-\u9fff]", text))


def fix_text(text):
    """纯字符串入口：按固定流水线修复一段文本。

    顺序不能随意调整：省略号必须早于单句号，破折号必须早于普通替换，成对
    引号放在最后处理。引号算法按出现次序交替左右引号，并不做语法或嵌套分析；
    若引号总数为奇数，最后一个会按左引号处理。
    """
    if not text:
        return text

    # 保护特殊模式
    result, protected = _protect_special_patterns(text)

    # ===== 第一步：处理省略号（必须在句号之前）=====
    result = re.sub(r"\.{2,}", "……", result)
    result = re.sub(r"。{2,}", "……", result)

    # ===== 第二步：处理破折号 =====
    result = re.sub(r"--+", "——", result)
    result = re.sub(r"—(?!—)", "——", result)

    # ===== 第三步：基本标点替换（只在有中文的文本中）=====
    # has_chinese 是整段开关：一旦段内含中文，映射表中的半角标点都会转换；
    # URL、时间等例外已在上面用占位符保护。
    if has_chinese(result):
        for en, cn in REPLACEMENTS.items():
            result = result.replace(en, cn)

    # ===== 第四步：逗号特殊处理 =====
    result = re.sub(r"([\u4e00-\u9fff]),", r"\1，", result)
    result = re.sub(r",([\u4e00-\u9fff])", r"，\1", result)

    # ===== 第五步：句号特殊处理 =====
    result = re.sub(r"([\u4e00-\u9fff])\.(\s|$)", r"\1。\2", result)

    # ===== 第六步：双引号处理 =====
    # 需要处理的双引号字符（使用Unicode确保正确）
    double_quote_chars = [
        '"',        # U+0022 ASCII直引号
        '\u201c',   # U+201C 左双引号 "
        '\u201d',   # U+201D 右双引号 "
        '\u201e',   # U+201E 双低引号 „
        '\u201f',   # U+201F 双高反引号 ‟
        '\u300c',   # U+300C 日文左引号 「
        '\u300d',   # U+300D 日文右引号 」
    ]

    # 统一替换为临时标记
    temp_result = result
    for q in double_quote_chars:
        temp_result = temp_result.replace(q, "\x00")

    # 配对处理
    if "\x00" in temp_result:
        chars = list(temp_result)
        quote_count = 0
        for i, c in enumerate(chars):
            if c == "\x00":
                # 偶数位置用左引号，奇数位置用右引号
                if quote_count % 2 == 0:
                    chars[i] = LEFT_DOUBLE_QUOTE  # "
                else:
                    chars[i] = RIGHT_DOUBLE_QUOTE  # "
                quote_count += 1
        result = "".join(chars)

    # ===== 第七步：单引号处理 =====
    single_quote_chars = [
        "'",        # U+0027 ASCII单引号
        '\u2018',   # U+2018 左单引号 '
        '\u2019',   # U+2019 右单引号 '
        '\u201a',   # U+201A 单低引号 ‚
        '\u201b',   # U+201B 单高反引号 ‛
    ]

    temp_result = result
    for q in single_quote_chars:
        temp_result = temp_result.replace(q, "\x01")

    if "\x01" in temp_result:
        chars = list(temp_result)
        quote_count = 0
        for i, c in enumerate(chars):
            if c == "\x01":
                # 偶数位置用左引号，奇数位置用右引号
                if quote_count % 2 == 0:
                    chars[i] = LEFT_SINGLE_QUOTE  # '
                else:
                    chars[i] = RIGHT_SINGLE_QUOTE  # '
                quote_count += 1
        result = "".join(chars)

    # 恢复被保护的内容
    result = _restore_protected(result, protected)
    return result


def _fix_simple_punctuation(text):
    """只做不依赖跨 run 上下文的标点替换，故意保留引号不动。

    文档入口会对每个 run 调用本函数，尽可能保留原有格式边界；引号随后拼接
    整段再配对，因为左、右引号可能被 Word 拆在不同 run 中。
    """
    if not text:
        return text

    # 保护特殊模式
    result, protected = _protect_special_patterns(text)

    # 省略号（必须在句号之前）
    result = re.sub(r"\.{2,}", "……", result)
    result = re.sub(r"。{2,}", "……", result)

    # 破折号
    result = re.sub(r"--+", "——", result)
    result = re.sub(r"—(?!—)", "——", result)

    # 基本替换（只在有中文时）
    if has_chinese(result):
        for en, cn in REPLACEMENTS.items():
            result = result.replace(en, cn)

    # 逗号
    result = re.sub(r"([\u4e00-\u9fff]),", r"\1，", result)
    result = re.sub(r",([\u4e00-\u9fff])", r"，\1", result)

    # 句号
    result = re.sub(r"([\u4e00-\u9fff])\.(\s|$)", r"\1。\2", result)

    # 恢复被保护的内容
    result = _restore_protected(result, protected)
    return result


def _fix_quotes_whole_text(text):
    """按出现次序把整段中的单双引号交替转换为左、右中文弯引号。"""
    result = text

    # 双引号
    double_quote_chars = ['"', "\u201c", "\u201d", "\u201e", "\u201f", "\u300c", "\u300d"]
    temp = result
    for q in double_quote_chars:
        temp = temp.replace(q, "\x00")

    if "\x00" in temp:
        chars = list(temp)
        quote_idx = 0
        for i, c in enumerate(chars):
            if c == "\x00":
                chars[i] = LEFT_DOUBLE_QUOTE if quote_idx % 2 == 0 else RIGHT_DOUBLE_QUOTE
                quote_idx += 1
        result = "".join(chars)

    # 单引号
    single_quote_chars = ["'", "\u2018", "\u2019", "\u201a", "\u201b"]
    temp = result
    for q in single_quote_chars:
        temp = temp.replace(q, "\x01")

    if "\x01" in temp:
        chars = list(temp)
        quote_idx = 0
        for i, c in enumerate(chars):
            if c == "\x01":
                chars[i] = LEFT_SINGLE_QUOTE if quote_idx % 2 == 0 else RIGHT_SINGLE_QUOTE
                quote_idx += 1
        result = "".join(chars)

    return result


def _redistribute_text_to_runs(runs, new_full_text):
    """按原长度边界把整段新文本分回各个 run，尽量保留原格式。

    普通的全角/半角替换不会改变字符数，因此可以安全地沿用原 run 边界。
    只有省略号等极少数规则会改变长度；这时已无法唯一推断每个新字符原来属于
    哪个 run，只能把结果放入首个 run。需要保留复杂 XML 的场景应使用下面的
    ``_split_run_around_characters``，而不是调用本函数重建内容。
    """
    run_lengths = [len(run.text) for run in runs]

    # 如果长度一致（只是字符替换，没有增删），直接按原长度切分
    total_original = sum(run_lengths)
    if len(new_full_text) == total_original:
        pos = 0
        for i, run in enumerate(runs):
            run.text = new_full_text[pos:pos + run_lengths[i]]
            pos += run_lengths[i]
    else:
        # 长度变了（极少情况），回退到旧方案：全塞第一个 run
        runs[0].text = new_full_text
        for run in runs[1:]:
            run.text = ""


def _run_font_values(run):
    """读取 run 上所有显式字体名，用于判断特殊字符是否已经规范化。"""
    values = []
    if run.font.name:
        values.append(run.font.name)
    rpr = run._element.rPr
    if rpr is not None and rpr.rFonts is not None:
        for attr in ("eastAsia", "ascii", "hAnsi", "cs"):
            value = rpr.rFonts.get(qn("w:" + attr))
            if value:
                values.append(value)
    return values


def _run_east_asian_font(run):
    """读取 OOXML 的 ``w:eastAsia`` 字体槽；未显式设置时返回 ``None``。"""
    rpr = run._element.rPr
    if rpr is None or rpr.rFonts is None:
        return None
    return rpr.rFonts.get(qn("w:eastAsia"))


def _set_run_font_slots(run, font_name, *, east_asian_hint=False):
    """为“只含目标标点”的 run 写入全部字体槽。

    ``python-docx`` 的 ``run.font.name`` 主要覆盖西文字体，并不足以决定中文
    字符实际显示的字体。Word 会在四个 OOXML 字体槽之间选择：

    - ``eastAsia``：中日韩字符；
    - ``ascii``：ASCII 英文和数字；
    - ``hAnsi``：高位 ANSI 字符，弯引号经常被 Word 归到这里；
    - ``cs``：复杂文字脚本。

    中文弯引号即使视觉上是中文标点，也可能因为 ``hAnsi`` 槽而显示成 Times
    New Roman，所以还要写入 ``w:hint="eastAsia"``。调用本函数前，目标标点
    已被拆成独立 run，因此覆盖四个槽不会影响相邻英文和数字的西文字体。
    """
    run.font.name = font_name
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("eastAsia", "ascii", "hAnsi", "cs"):
        rfonts.set(qn("w:" + attr), font_name)
    if east_asian_hint:
        rfonts.set(qn("w:hint"), "eastAsia")


def _run_has_only_font(run, font_name):
    rpr = run._element.rPr
    if rpr is None or rpr.rFonts is None:
        return False
    return all(
        rpr.rFonts.get(qn("w:" + attr)) == font_name
        for attr in ("eastAsia", "ascii", "hAnsi", "cs")
    )


def _iter_run_elements(container):
    """按文档顺序遍历段落内的 run，包括被包装的 run。

    ``Paragraph.runs`` 只返回段落的直接子 run；超链接 ``w:hyperlink``、内容
    控件 ``w:sdt`` 和修订 ``w:ins`` 等节点内部的 run 会被漏掉。本函数递归这些
    包装节点，但遇到嵌套 ``w:p`` 就停止，因为文本框内的新段落属于另一条文档
    story，不能算作外层段落的一部分。
    """
    for child in container:
        if child.tag == qn("w:p"):
            continue
        if child.tag == qn("w:r"):
            yield child
            continue
        yield from _iter_run_elements(child)


def iter_paragraph_runs(para):
    """返回一个段落的全部直接/包装 run，并转换成 python-docx ``Run``。"""
    return [Run(run_element, para) for run_element in _iter_run_elements(para._p)]


def _split_character_parts(text, target_chars):
    """把文本切成“目标字符”和“普通文本”交替出现的非空片段。"""
    pattern = "(" + "|".join(re.escape(char) for char in sorted(target_chars)) + ")"
    return [part for part in re.split(pattern, text) if part]


def _clone_run_shell(run_element):
    """复制 run 的属性和 ``w:rPr``，暂不复制实际内容节点。"""
    new_run = OxmlElement("w:r")
    for attribute, value in run_element.attrib.items():
        new_run.set(attribute, value)
    rpr = run_element.find(qn("w:rPr"))
    if rpr is not None:
        new_run.append(deepcopy(rpr))
    return new_run


def _clone_text_element(text_element, text):
    """复制一个 ``w:t`` 并替换文本，同时维护 XML 空格保留标记。"""
    clone = deepcopy(text_element)
    clone.text = text
    if text[:1].isspace() or text[-1:].isspace():
        clone.set(_XML_SPACE, "preserve")
    else:
        clone.attrib.pop(_XML_SPACE, None)
    return clone


def _split_run_around_characters(run, target_chars, font_name, *, east_asian_hint=False):
    """把目标字符隔离成独立 run，设置字体，并保留原 run 的其他 XML。

    算法要点：

    1. 如果 run 全部由目标字符组成，直接改字体即可；
    2. 否则逐个处理原 run 的子节点；只有 ``w:t`` 文本节点会被切分；
    3. 图片、域代码、制表符、换行等非文本子节点整体深拷贝，绝不转成字符串；
    4. 新 run 继承原 ``w:rPr``，只有目标字符片段再覆盖字体；
    5. 所有新 run 就位后才删除原 run，保持原来的文档顺序。

    返回值表示 XML 是否发生改变，便于上层统计和保持幂等性。
    """
    text = run.text
    if not text or not any(char in target_chars for char in text):
        return False

    if all(char in target_chars for char in text):
        already_formatted = _run_has_only_font(run, font_name)
        if east_asian_hint:
            rpr = run._element.rPr
            rfonts = rpr.rFonts if rpr is not None else None
            already_formatted = (
                already_formatted
                and rfonts is not None
                and rfonts.get(qn("w:hint")) == "eastAsia"
            )
        if already_formatted:
            return False
        _set_run_font_slots(run, font_name, east_asian_hint=east_asian_hint)
        return True

    original_r = run._r
    parent = original_r.getparent()
    if parent is None:
        return False

    # 先在内存中组装替代 run；全部成功后再一次性替换，避免处理中途留下
    # 半成品 XML。
    split_runs = []
    for child in original_r:
        if child.tag == qn("w:rPr"):
            continue
        if child.tag == qn("w:t") and any(char in target_chars for char in (child.text or "")):
            for part in _split_character_parts(child.text or "", target_chars):
                new_r = _clone_run_shell(original_r)
                new_r.append(_clone_text_element(child, part))
                split_runs.append((new_r, all(char in target_chars for char in part)))
        else:
            new_r = _clone_run_shell(original_r)
            new_r.append(deepcopy(child))
            split_runs.append((new_r, False))

    if not split_runs:
        return False

    insertion_index = parent.index(original_r)
    for offset, (new_r, is_target) in enumerate(split_runs):
        parent.insert(insertion_index + offset, new_r)
        if is_target:
            _set_run_font_slots(
                Run(new_r, run._parent),
                font_name,
                east_asian_hint=east_asian_hint,
            )
    parent.remove(original_r)

    return True


def _contextual_east_asian_font(runs, index):
    """推断某个 run 所在位置应使用的中文字体。

    优先读取 run 自己的 ``eastAsia`` 槽；若该 run 只含标点而没有显式字体，
    就从左右两侧按距离由近到远查找。这样独立标点 run 仍会继承其所在标题、
    正文或各级标题的中文字体，而不是使用一个全局固定字体。
    """
    own_font = _run_east_asian_font(runs[index])
    if own_font:
        return own_font

    for distance in range(1, len(runs)):
        for candidate_index in (index - distance, index + distance):
            if 0 <= candidate_index < len(runs):
                candidate = _run_east_asian_font(runs[candidate_index])
                if candidate:
                    return candidate
    return None


def normalize_chinese_quote_fonts(para, font_cn=None):
    """让中文弯引号使用其所在位置的中文字体。

    先隔离引号、后改字体，确保相邻英文和阿拉伯数字仍使用段落的西文字体。
    格式化阶段会显式传入 ``font_cn``；仅修复标点时没有预设上下文，就从引号
    自身或最近 run 的 ``eastAsia`` 槽推断。
    """
    changed = False
    runs = iter_paragraph_runs(para)
    for index, run in enumerate(runs):
        if not any(char in CHINESE_QUOTE_CHARS for char in run.text):
            continue
        contextual_font = font_cn or _contextual_east_asian_font(runs, index)
        if contextual_font and _split_run_around_characters(
            run,
            CHINESE_QUOTE_CHARS,
            contextual_font,
            east_asian_hint=True,
        ):
            changed = True
    return changed


def middle_dot_run_has_required_font(run):
    """检查含 ``·`` 的 run 上已显式声明的字体名是否都属于允许字体。"""
    values = _run_font_values(run)
    return bool(values) and all(value in MIDDLE_DOT_ALLOWED_FONTS for value in values)


def _set_middle_dot_font(run):
    _set_run_font_slots(run, MIDDLE_DOT_FONT_CN, east_asian_hint=True)


def _split_run_around_middle_dot(run):
    return _split_run_around_characters(
        run,
        {MIDDLE_DOT},
        MIDDLE_DOT_FONT_CN,
        east_asian_hint=True,
    )


def normalize_middle_dot_fonts(para):
    """把每个 U+00B7 ``·`` 设为方正仿宋_GBK，并保留周围 run 内容。"""
    changed = False
    for run in iter_paragraph_runs(para):
        if _split_run_around_middle_dot(run):
            changed = True
    return changed


def _process_spaces_text(text, mode):
    """按空格策略返回新文本。

    ``keep_all`` 完全保留；``remove_all`` 删除半角与全角空格；
    ``keep_en_boundary`` 删除汉字之间的空格，并把“汉字 ↔ ASCII 英文/数字”
    边界统一成一个空格。中文范围覆盖基本区、扩展 A 和兼容区。
    """
    if mode == 'keep_all' or not text:
        return text
    if mode == 'remove_all':
        # 删除所有半角空格和全角空格
        return text.replace('\u3000', '').replace(' ', '')
    if mode == 'keep_en_boundary':
        CN = r'\u4e00-\u9fff\u3400-\u4dbf\uf900-\ufaff'
        EN = r'[A-Za-z0-9]'
        CN_cls = f'[{CN}]'

        # 第一步：删除纯中文之间的空格（不涉及英文/数字边界）
        text = re.sub(f'(?<={CN_cls}) +(?={CN_cls})', '', text)

        # 第二步：将中文与英文/数字边界处的空格（0个或多个）统一替换为恰好1个
        # 中文 → 英文/数字
        text = re.sub(f'(?<={CN_cls}) *(?={EN})', ' ', text)
        # 英文/数字 → 中文
        text = re.sub(f'(?<={EN}) *(?={CN_cls})', ' ', text)

        # 第三步：修正因第二步可能在段落首尾产生的多余前导/尾随空格
        # 段落开头的英文/数字不应有前导空格
        text = re.sub(r'^ +', '', text)
        # 段落结尾的英文/数字不应有尾随空格
        text = re.sub(r' +$', '', text)

        return text
    return text


def process_spaces(para, mode='remove_all'):
    """处理段落内空格，返回是否有改动。

    这里使用 python-docx 的直接 ``para.runs``；超链接等包装节点中的文本不会
    被纳入空格重分配。长度变化时还会触发 ``_redistribute_text_to_runs`` 的
    有损回退，维护者应避免在此加入会大幅改变字符数的规则。
    """
    if mode == 'keep_all':
        return False
    full_text = para.text
    if not full_text.strip():
        return False
    new_text = _process_spaces_text(full_text, mode)
    if new_text == full_text:
        return False
    _redistribute_text_to_runs(para.runs, new_text)
    return True


def process_paragraph(para, space_mode='remove_all'):
    """完成一个段落的标点、空格和特殊字符字体处理。

    执行顺序是有意设计的：先逐 run 做不依赖上下文的替换，以保留各 run 格式；
    再把整段文本拼起来配对引号；随后处理空格；最后才隔离引号和中点并设置
    字体。字体归一化放在末尾，可避免前面的文本重分配再次覆盖新建的 run。
    """
    full_text = para.text
    if not full_text.strip():
        return False

    runs = para.runs
    if not runs:
        return False

    changed = False

    # 第一步：逐 run 做简单标点替换（保留所有格式）
    for run in runs:
        original = run.text
        fixed = _fix_simple_punctuation(original)
        if fixed != original:
            run.text = fixed
            changed = True

    # 第二步：引号配对需要整段处理
    full_after_simple = para.text
    full_after_quotes = _fix_quotes_whole_text(full_after_simple)

    if full_after_quotes != full_after_simple:
        _redistribute_text_to_runs(runs, full_after_quotes)
        changed = True

    # 空格处理
    if process_spaces(para, space_mode):
        changed = True

    if normalize_chinese_quote_fonts(para):
        changed = True

    if normalize_middle_dot_fonts(para):
        changed = True

    return changed


def process_document(input_path, output_path):
    """读取、处理并保存一篇 DOCX，供命令行脚本直接调用。

    GUI 的“智能处理”会调用更完整的调度链；这里保留为独立标点修复入口。
    ``Document.paragraphs`` 不包含表格单元格段落，因此正文和顶层表格需要分别
    遍历。字符替换不递归页眉、页脚和文本框；最后的中文换行规则覆盖范围更广。
    """
    print(f"Reading: {input_path}")
    doc = Document(input_path)

    changes = 0

    # 处理段落
    for i, para in enumerate(doc.paragraphs):
        if process_paragraph(para):
            changes += 1
            preview = para.text[:50] + "..." if len(para.text) > 50 else para.text
            print(f"  Para {i + 1}: {preview}")

    # 处理表格
    table_changes = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if process_paragraph(para):
                        table_changes += 1

    if table_changes > 0:
        print(f"  Tables: {table_changes} cells fixed")

    apply_chinese_line_break_rules(doc)

    print()
    print(f"Total: {changes} paragraphs + {table_changes} table cells fixed")
    doc.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python punctuation.py input.docx output.docx")
        sys.exit(1)

    process_document(sys.argv[1], sys.argv[2])

# 测试用例（验证 fix_text）
# assert fix_text('会议时间:上午9:30至下午14:30,请准时参加.') == '会议时间：上午9:30至下午14:30，请准时参加。'
# assert fix_text('请于14:30前将材料发送至 report@gov.cn,逾期不候.') == '请于14:30前将材料发送至 report@gov.cn，逾期不候。'
# assert fix_text('详情请访问 https://www.example.com:8080/path 了解.') == '详情请访问 https://www.example.com:8080/path 了解。'
# assert fix_text('参照 ISO 9001:2015 执行.') == '参照 ISO 9001:2015 执行。'
# assert fix_text('每日9:00前完成巡检.') == '每日9:00前完成巡检。'
# assert '还需要改进' in fix_text('他说"这个方案"不错，但是"还需要改进')  # 奇数引号照常替换
